from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from inventura.core.documents import split_by_rack
from inventura.core.locations import parse_location
from inventura.core.stock import StockRow
from inventura.io.count_sheet_reader import CountSheetError, read_count_sheet
from inventura.io.count_sheets import SheetOptions, write_count_sheets_xlsx


def stock(lokacija: str, sifra: str, sarza: str | None = None) -> StockRow:
    return StockRow(
        0, sifra, "Material", "kos", parse_location(lokacija), sarza, Decimal(5), Decimal(1)
    )


def written_sheets(book_quantities: bool = False) -> bytes:
    documents = split_by_rack(
        [stock("B6-1-1", "0000001"), stock("B6-1-2", "0000002", "L1"), stock("K2-01-01", "0000003")]
    )
    buffer = BytesIO()
    options = SheetOptions(date(2026, 10, 5), len(documents), show_book_quantity=book_quantities)
    write_count_sheets_xlsx(documents, buffer, options)
    return buffer.getvalue()


def filled(data: bytes, sheet: str, values: dict[str, object]) -> bytes:
    workbook = load_workbook(BytesIO(data))
    for cell, value in values.items():
        workbook[sheet][cell] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_reads_the_rack_worksheet_with_filled_quantities() -> None:
    data = filled(
        written_sheets(), "B6", {"G6": 4, "G7": "2,5", "B8": "B6-2-2", "C8": "9", "G8": 1}
    )
    rows = read_count_sheet(data, "B6")
    assert [(r.row_number, r.lokacija, r.sifra, r.sarza, r.kolicina) for r in rows] == [
        (6, "B6-1-1", "0000001", None, 4),
        (7, "B6-1-2", "0000002", "L1", "2,5"),
        (8, "B6-2-2", "9", None, 1),
    ]


def test_count_column_is_found_when_book_quantities_are_printed() -> None:
    data = filled(written_sheets(book_quantities=True), "K2", {"H6": 7})
    (only,) = read_count_sheet(data, "K2")
    assert (only.sifra, only.kolicina) == ("0000003", 7)


def test_single_worksheet_with_another_name_is_used() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Lokacija", "Šifra", "Prešteta količina"])
    sheet.append(["B6-1-1", "0000001", 3])
    buffer = BytesIO()
    workbook.save(buffer)
    (only,) = read_count_sheet(buffer.getvalue(), "B6")
    assert (only.row_number, only.kolicina, only.sarza) == (2, 3, None)


def test_workbook_without_the_rack() -> None:
    with pytest.raises(CountSheetError, match="ni lista z imenom 'B9'"):
        read_count_sheet(written_sheets(), "B9")


def test_missing_columns() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Šifra", "Prešteta količina"])
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(CountSheetError, match="manjkajo stolpci: Lokacija"):
        read_count_sheet(buffer.getvalue(), "B6")


def test_no_header_row() -> None:
    workbook = Workbook()
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(CountSheetError, match="ni vrstice z naslovom stolpca"):
        read_count_sheet(buffer.getvalue(), "B6")


def test_not_an_excel_file() -> None:
    with pytest.raises(CountSheetError, match="ni mogoče prebrati"):
        read_count_sheet(b"Lokacija;Sifra\n", "B6")
