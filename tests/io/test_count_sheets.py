from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from inventura.core.documents import RackDocument, split_by_rack
from inventura.core.locations import parse_location
from inventura.core.stock import StockRow
from inventura.io.count_sheets import (
    HEADER_ROW,
    SheetOptions,
    render_count_sheets_html,
    write_count_sheets_html,
    write_count_sheets_xlsx,
)

OPTIONS = SheetOptions(created=date(2026, 9, 28), total_documents=2)


def row(
    lokacija: str, sifra: str, unit: str = "kos", qty: str = "5", opis: str = "Vijak"
) -> StockRow:
    return StockRow(
        row_number=0,
        sifra=sifra,
        opis=opis,
        merska_enota=unit,
        lokacija=parse_location(lokacija),
        sarza=None,
        kolicina=Decimal(qty),
        cena_na_enoto=Decimal("1.00"),
    )


@pytest.fixture
def documents() -> list[RackDocument]:
    return split_by_rack(
        [
            row("B10-1-1", "0000003"),
            row("B2-1-1", "0000001", "m", "12.5", opis="Kabel <NYM> & co"),
            row("B2-1-2", "0000002", qty="1234"),
            row("B2-2-1", "0000004"),
        ]
    )


def sheet_rows(sheet: Worksheet) -> list[tuple[object, ...]]:
    return list(sheet.iter_rows(min_row=HEADER_ROW, values_only=True))


def test_xlsx_has_one_sheet_per_rack(documents: list[RackDocument], tmp_path: Path) -> None:
    path = tmp_path / "sheets.xlsx"
    write_count_sheets_xlsx(documents, path, OPTIONS)
    workbook = load_workbook(path)

    assert workbook.sheetnames == ["B2", "B10"]
    sheet = workbook["B2"]
    assert sheet["A1"].value == "Popisni list – regal B2"
    assert sheet["A2"].value == "Dokument 1 od 2 · Datum: 28. 9. 2026 · Postavk: 3 · Lokacij: 3"
    assert sheet_rows(sheet) == [
        ("Zap. št.", "Lokacija", "Šifra", "Opis", "Šarža", "ME", "Prešteta količina", "Opomba"),
        (1, "B2-1-1", "0000001", "Kabel <NYM> & co", None, "m", None, None),
        (2, "B2-1-2", "0000002", "Vijak", None, "kos", None, None),
        (3, "B2-2-1", "0000004", "Vijak", None, "kos", None, None),
    ]


def test_xlsx_count_column_is_set_up_for_entry(
    documents: list[RackDocument], tmp_path: Path
) -> None:
    path = tmp_path / "sheets.xlsx"
    write_count_sheets_xlsx(documents, path, OPTIONS)
    sheet = load_workbook(path)["B2"]

    assert [sheet[f"G{r}"].number_format for r in (6, 7)] == ["#,##0.0", "#,##0"]
    validations = {dv.type: str(dv.sqref) for dv in sheet.data_validations.dataValidation}
    assert validations == {"decimal": "G6", "whole": "G7 G8"}
    assert sheet.print_title_rows == "$5:$5"
    assert sheet.freeze_panes == "A6"
    assert sheet.page_setup.orientation == "portrait"
    assert sheet.oddFooter.right.text == "Stran &P od &N"
    # A new level starts on row 8 (B2-2-1) and gets a heavier top border.
    assert sheet["A8"].border.top.style == "medium"
    assert sheet["A7"].border.top.style == "thin"


def test_xlsx_with_book_quantities(documents: list[RackDocument], tmp_path: Path) -> None:
    path = tmp_path / "sheets.xlsx"
    options = SheetOptions(created=date(2026, 9, 28), total_documents=2, show_book_quantity=True)
    write_count_sheets_xlsx(documents, path, options)
    sheet = load_workbook(path)["B2"]

    header = sheet_rows(sheet)[0]
    assert header[6:8] == ("Knjižna količina", "Prešteta količina")
    assert [sheet[f"G{r}"].value for r in (6, 7)] == [12.5, 1234]
    assert sheet.page_setup.orientation == "landscape"
    assert {dv.type: str(dv.sqref) for dv in sheet.data_validations.dataValidation} == {
        "decimal": "H6",
        "whole": "H7 H8",
    }


def test_xlsx_requires_documents(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        write_count_sheets_xlsx([], tmp_path / "x.xlsx", OPTIONS)


def test_html_lists_documents_in_order(documents: list[RackDocument]) -> None:
    html = render_count_sheets_html(documents, OPTIONS)
    assert html.index("Popisni list – regal B2") < html.index("Popisni list – regal B10")
    assert html.count('<section class="sheet">') == 2
    assert html.count('<td class="count"></td>') == 4
    assert html.count('class="new-level"') == 1
    assert "Dokument 2 od 2" in html
    assert "Knjižna količina" not in html


def test_html_escapes_text(documents: list[RackDocument]) -> None:
    html = render_count_sheets_html(documents, OPTIONS)
    assert "Kabel &lt;NYM&gt; &amp; co" in html
    assert "<NYM>" not in html


def test_html_book_quantities_use_unit_decimals(documents: list[RackDocument]) -> None:
    options = SheetOptions(created=date(2026, 9, 28), total_documents=2, show_book_quantity=True)
    html = render_count_sheets_html(documents, options)
    assert '<td class="book">12,5</td>' in html
    assert '<td class="book">1.234</td>' in html
    assert '<section class="sheet wide">' in html


def test_html_has_print_styles(documents: list[RackDocument], tmp_path: Path) -> None:
    path = tmp_path / "sheets.html"
    write_count_sheets_html(documents, path, OPTIONS)
    html = path.read_text(encoding="utf-8")
    assert "@media print" in html
    assert "break-after: page" in html
    assert "size: A4" in html
