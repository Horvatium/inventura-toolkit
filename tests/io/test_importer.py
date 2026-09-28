import csv
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from inventura.core.stock import StockField
from inventura.io.column_mapping import ColumnMapping, ColumnMappingError
from inventura.io.importer import UnsupportedFileError, import_stock_file, write_error_report

F = StockField


def test_csv_with_errors_reports_them_by_row(mapping: ColumnMapping, fixtures_dir: Path) -> None:
    report = import_stock_file(fixtures_dir / "stock_errors.csv", mapping)
    result = report.result

    assert report.data_rows == 11
    assert [row.row_number for row in result.rows] == [2, 12]
    assert result.rows[1].kolicina == Decimal("1250.500")
    assert result.skipped_blank_rows == 2
    assert [(e.row_number, e.field) for e in result.errors] == [
        (4, F.OPIS),
        (5, F.KOLICINA),
        (6, F.KOLICINA),
        (7, F.KOLICINA),
        (8, None),
        (9, F.MERSKA_ENOTA),
        (10, F.KOLICINA),
    ]


def test_csv_keeps_leading_zeros_and_text(mapping: ColumnMapping, fixtures_dir: Path) -> None:
    row = import_stock_file(fixtures_dir / "stock_errors.csv", mapping).result.rows[0]
    assert row.sifra == "0000101"
    assert row.opis == "Vijak šestrobi M8x40"


def test_csv_with_other_header_names(mapping: ColumnMapping, fixtures_dir: Path) -> None:
    report = import_stock_file(fixtures_dir / "stock_english_headers.csv", mapping)
    assert report.result.is_valid
    assert F.SARZA not in report.columns
    assert [(r.sifra, r.lokacija, r.kolicina) for r in report.result.rows] == [
        ("0000201", "B3-2-4", Decimal("12.000")),
        ("0000202", "K2-03-11", Decimal("18.250")),
    ]


def test_xlsx_numeric_and_text_cells(mapping: ColumnMapping, tmp_path: Path) -> None:
    path = tmp_path / "stock.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Šifra materiala", "Opis materiala", "ME", "Lokacija", "Zaloga", "Cena na enoto"])
    sheet.append(["0000301", "Olje HLP 46", "l", "K1-02-03", Decimal("208.125"), 3.9])
    sheet.append([None, None, None, None, None, None])
    sheet.append([302, "Filter F120", "kos", "B4-1-2", 7, "1.234,50"])
    sheet.append(["0000303", "Rokavice nitril", "par", "B5-3-1", 0.1, "x"])
    workbook.save(path)

    result = import_stock_file(path, mapping).result

    assert [(r.row_number, r.sifra, r.kolicina, r.cena_na_enoto) for r in result.rows] == [
        (2, "0000301", Decimal("208.125"), Decimal("3.90")),
        (4, "302", Decimal("7.000"), Decimal("1234.50")),
    ]
    assert result.skipped_blank_rows == 1
    assert [(e.row_number, e.field) for e in result.errors] == [(5, F.CENA_NA_ENOTO)]


def test_xlsx_sheet_by_name(mapping: ColumnMapping, tmp_path: Path) -> None:
    path = tmp_path / "stock.xlsx"
    workbook = Workbook()
    workbook.create_sheet("Zaloga")
    workbook["Zaloga"].append(["Material", "Opis", "ME", "Lokacija", "Zaloga", "Cena"])
    workbook["Zaloga"].append(["1", "Vijak", "kos", "B1-1-1", 1, 1])
    workbook.save(path)

    by_name = mapping.model_copy(
        update={"xlsx": mapping.xlsx.model_copy(update={"sheet": "Zaloga"})}
    )
    assert len(import_stock_file(path, by_name).result.rows) == 1


def test_missing_columns_raise(mapping: ColumnMapping, tmp_path: Path) -> None:
    path = tmp_path / "stock.csv"
    path.write_text("Šifra materiala;Opis materiala\n1;Vijak\n", encoding="utf-8")
    with pytest.raises(ColumnMappingError, match="missing columns"):
        import_stock_file(path, mapping)


def test_unsupported_file_type(mapping: ColumnMapping, tmp_path: Path) -> None:
    path = tmp_path / "stock.txt"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedFileError, match=r"\.txt"):
        import_stock_file(path, mapping)


def test_error_report_csv(mapping: ColumnMapping, fixtures_dir: Path, tmp_path: Path) -> None:
    errors = import_stock_file(fixtures_dir / "stock_errors.csv", mapping).result.errors
    path = tmp_path / "errors.csv"
    write_error_report(errors, path)

    with path.open(encoding="utf-8-sig", newline="") as file:
        lines = list(csv.reader(file, delimiter=";"))
    assert lines[0] == ["row", "field", "value", "message"]
    assert lines[1] == ["4", "opis", "", "value is required"]
    assert len(lines) == len(errors) + 1
