from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from inventura.core.variance import VarianceRules, compute_variance
from inventura.io.reports import (
    ReportHeader,
    ReportLine,
    ReportRack,
    rules_text,
    write_variance_report,
)
from inventura.io.variance_rules import load_variance_rules

REPO_CONFIG = Path(__file__).resolve().parents[2] / "config" / "variance_rules.yaml"
RULES = VarianceRules()
D = Decimal


def test_repository_rules_are_the_business_rules() -> None:
    rules = load_variance_rules(REPO_CONFIG)
    assert rules == VarianceRules(D("50.00"), D("5"), D("500.00"), 2)
    assert isinstance(rules.min_value, Decimal)
    assert str(rules.min_value) == "50.00"  # read as text, never through float


def test_invalid_rules_file(tmp_path: Path) -> None:
    path = tmp_path / "rules.yaml"
    path.write_text("min_value: -1\nmin_percent: 5\nalways_value: 500\nmax_rounds: 2\n")
    with pytest.raises(ValidationError):
        load_variance_rules(path)
    path.write_text("min_value: 50\nmin_percent: 5\nalways_value: 500\nmax_rounds: 2\nx: 1\n")
    with pytest.raises(ValidationError):
        load_variance_rules(path)


def test_rules_text() -> None:
    assert rules_text(RULES) == (
        "Ponovno štetje: (|vrednost razlike| ≥ 50,00 € in |odstopanje| ≥ 5 %) "
        "ali |vrednost razlike| ≥ 500,00 €; največ 2 kroga štetja."
    )


def line(
    lokacija: str, book: str, counted: str | None, price: str = "10.00", **kw: object
) -> ReportLine:
    variance = None if counted is None else compute_variance(D(book), D(counted), D(price), RULES)
    return ReportLine(
        lokacija=lokacija,
        sifra=str(kw.get("sifra", "0000001")),
        opis="Material",
        sarza=None,
        merska_enota=str(kw.get("unit", "kos")),
        krog=int(kw.get("krog", 1)),  # type: ignore[call-overload]
        najdeno=bool(kw.get("najdeno", False)),
        variance=variance,
    )


def report(racks: list[ReportRack]) -> BytesIO:
    buffer = BytesIO()
    header = ReportHeader(snapshot="Uvoz 1 · zaloga.csv", created=date(2026, 10, 5), rules=RULES)
    write_variance_report(racks, header, buffer)
    buffer.seek(0)
    return buffer


def test_report_sheets_and_summary() -> None:
    racks = [
        ReportRack(
            1,
            "B2",
            "Zaključen",
            [
                line("B2-1-1", "10", "4", "20.00"),  # -120.00, recount
                line("B2-1-2", "5", "5"),  # no difference
                line("B2-1-3", "0", "2", "3.00", najdeno=True),  # +6.00 found
            ],
        ),
        ReportRack(2, "B10", "V štetju", [line("B10-1-1", "3", None, sifra="0000009")]),
    ]
    workbook = load_workbook(report(racks))
    assert workbook.sheetnames == ["Povzetek", "Razlike", "Neprešteto"]

    summary = workbook["Povzetek"]
    assert summary["A1"].value == "Inventura – poročilo o razlikah"
    rows = list(summary.iter_rows(min_row=6, max_row=8, values_only=True))
    assert rows[0] == (1, "B2", "Zaključen", 3, 3, 0, 2, 1, 6, -120, -114, 126)
    assert rows[1][:8] == (2, "B10", "V štetju", 1, 0, 1, 0, 0)
    assert rows[2][1] == "Skupaj"
    assert rows[2][3:] == (4, 3, 1, 2, 1, 6, -120, -114, 126)


def test_report_variance_rows() -> None:
    racks = [
        ReportRack(
            1,
            "B2",
            "Zaključen",
            [
                line("B2-1-1", "10", "4", "20.00", krog=2),
                line("B2-1-2", "5", "5"),
                line("B2-1-3", "0", "2.5", "3.00", unit="m", najdeno=True),
            ],
        )
    ]
    sheet = load_workbook(report(racks))["Razlike"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2  # the item without a difference is left out
    assert rows[0][:4] == ("B2", "B2-1-1", "0000001", "Material")
    assert rows[0][6:] == (10, 4, -6, 20, -120, -0.6, 2, "da", None)
    assert rows[1][6:] == (0, 2.5, 2.5, 3, 7.5, None, 1, "ne", "da")
    assert sheet["L2"].number_format == "0.0%"
    assert sheet.auto_filter.ref == "A1:O3"
    ranges = {str(rule.sqref) for rule in sheet.conditional_formatting}
    assert {"K2:K3", "A2:O3", "N2:N3"} <= ranges


def test_report_without_variances_or_uncounted_items() -> None:
    workbook = load_workbook(report([ReportRack(1, "B2", "Zaključen", [line("B2-1-1", "1", "1")])]))
    assert workbook["Razlike"]["A2"].value == "Ni razlik med knjižnim in preštetim stanjem."
    assert workbook["Neprešteto"]["A2"].value == "Vse postavke so preštete."
