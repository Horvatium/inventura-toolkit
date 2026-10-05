from io import BytesIO
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from inventura.generator import generate_stock, write_csv


def item_id(document: dict[str, Any], sifra: str) -> int:
    found: int = next(i["id"] for i in document["postavke"] if i["sifra"] == sifra)
    return found


def test_index_lists_imports_and_documents(client: TestClient, document: dict[str, Any]) -> None:
    page = client.get("/")
    assert page.status_code == 200
    assert "stock.csv" in page.text
    assert f'href="/documents/{document["id"]}"' in page.text
    assert "0 / 3" in page.text


def test_import_from_the_page(client: TestClient, tmp_path: Any) -> None:
    path = tmp_path / "zaloga.csv"
    write_csv(generate_stock(materials=5, seed=2), path)

    response = client.post("/snapshots", files={"file": ("zaloga.csv", path.read_bytes())})
    assert response.status_code == 204
    assert response.headers["HX-Redirect"] == "/"

    bad = client.post("/snapshots", files={"file": ("x.csv", b"A;B\n1;2\n")})
    assert "Datoteke ni mogoče uvoziti" in bad.text


def test_start_count_from_the_page(client: TestClient) -> None:
    data = "Šifra materiala;Opis materiala;ME;Lokacija;Zaloga;Cena na enoto\n1;V;kos;B1-1-1;1;1\n"
    snapshot = client.post("/api/snapshots", files={"file": ("s.csv", data.encode())}).json()
    response = client.post(f"/snapshots/{snapshot['id']}/documents", follow_redirects=False)
    assert response.status_code == 303
    assert client.get("/api/documents", params={"snapshot_id": snapshot["id"]}).json()


def test_count_page_is_a_blind_count(client: TestClient, document: dict[str, Any]) -> None:
    page = client.get(f"/documents/{document['id']}")
    assert page.status_code == 200
    assert page.text.count('class="qty"') == 3
    assert "Regal B6" in page.text
    # Book quantities (120, 12,5 and 20) are not shown to the counter.
    for book in ("120", "12,5", ">20<", "Knjižna"):
        assert book not in page.text


def test_saving_a_count_returns_the_row_and_progress(
    client: TestClient, document: dict[str, Any]
) -> None:
    response = client.patch(
        f"/items/{item_id(document, '0000002')}", data={"kolicina": "12,5", "stevec": "Ana"}
    )
    assert response.status_code == 200
    assert 'class="counted"' in response.text
    assert 'value="12,5"' in response.text
    assert 'id="doc-progress" class="doc-progress" hx-swap-oob="true"' in response.text
    assert "1 / 3" in response.text
    saved = client.get(f"/api/documents/{document['id']}").json()["postavke"]
    assert [(i["presteta_kolicina"], i["stevec"]) for i in saved if i["sifra"] == "0000002"] == [
        ("12.500", "Ana")
    ]


def test_invalid_count_shows_the_error_in_the_row(
    client: TestClient, document: dict[str, Any]
) -> None:
    response = client.patch(f"/items/{item_id(document, '0000001')}", data={"kolicina": "1,5"})
    assert response.status_code == 200
    assert "has-error" in response.text
    assert 'value="1,5"' in response.text
    assert "must be a whole number" in response.text
    assert client.get(f"/api/documents/{document['id']}").json()["presteto"] == 0


def test_clearing_a_count(client: TestClient, document: dict[str, Any]) -> None:
    url = f"/items/{item_id(document, '0000001')}"
    client.patch(url, data={"kolicina": "0"})
    assert client.get(f"/api/documents/{document['id']}").json()["presteto"] == 1
    response = client.patch(url, data={"kolicina": ""})
    assert 'class="counted"' not in response.text
    assert client.get(f"/api/documents/{document['id']}").json()["presteto"] == 0


def test_found_goods_form(client: TestClient, document: dict[str, Any]) -> None:
    url = f"/documents/{document['id']}/found"
    form = {"lokacija": "B6-3-1", "sifra": "0000009", "sarza": "", "kolicina": "2"}

    added = client.post(url, data=form)
    assert "Dodano: 0000009 na B6-3-1" in added.text
    # The refreshed item list is sent out of band inside a template tag.
    assert '<template><tbody id="items" hx-swap-oob="true">' in added.text
    assert "najdeno" in added.text

    rejected = client.post(url, data=form | {"sifra": "7777777"})
    assert "not in the material master data" in rejected.text
    assert 'value="7777777"' in rejected.text
    assert "hx-swap-oob" not in rejected.text


def test_delete_found_goods_from_the_page(client: TestClient, document: dict[str, Any]) -> None:
    url = f"/documents/{document['id']}/found"
    client.post(url, data={"lokacija": "B6-3-1", "sifra": "0000009", "kolicina": "2"})
    found_id = next(
        i["id"]
        for i in client.get(f"/api/documents/{document['id']}").json()["postavke"]
        if i["najdeno"]
    )
    response = client.delete(f"/items/{found_id}")
    assert response.status_code == 200
    assert "0 / 3" in response.text


def test_upload_from_the_page(client: TestClient, document: dict[str, Any]) -> None:
    sheet = client.get(f"/api/documents/{document['id']}/count-sheet.xlsx").content
    workbook = load_workbook(BytesIO(sheet))
    workbook["B6"]["G6"] = 120
    buffer = BytesIO()
    workbook.save(buffer)

    url = f"/documents/{document['id']}/upload"
    response = client.post(url, files={"file": ("b6.xlsx", buffer.getvalue())})
    assert "Popisni list je uvožen." in response.text
    assert "Vpisane količine: 1" in response.text
    assert '<template><tbody id="items" hx-swap-oob="true">' in response.text

    bad = client.post(url, files={"file": ("b6.xlsx", b"nonsense")})
    assert "To ni popisni list tega regala" in bad.text


def test_import_with_row_errors_lists_them(client: TestClient, fixtures_dir: Any) -> None:
    data = (fixtures_dir / "stock_errors.csv").read_bytes()
    response = client.post("/snapshots", files={"file": ("stock_errors.csv", data)})
    assert "Datoteka ima napake, nič ni uvoženo." in response.text
    assert "Vrstic z napakami: 10" in response.text
    assert "value is required" in response.text


def test_upload_with_row_errors_lists_them(client: TestClient, document: dict[str, Any]) -> None:
    sheet = client.get(f"/api/documents/{document['id']}/count-sheet.xlsx").content
    workbook = load_workbook(BytesIO(sheet))
    workbook["B6"]["G6"] = 1.5
    buffer = BytesIO()
    workbook.save(buffer)
    url = f"/documents/{document['id']}/upload"
    response = client.post(url, files={"file": ("b6.xlsx", buffer.getvalue())})
    assert "Popisni list ima napake, nič ni shranjeno." in response.text
    assert "must be a whole number" in response.text
    assert "hx-swap-oob" not in response.text
