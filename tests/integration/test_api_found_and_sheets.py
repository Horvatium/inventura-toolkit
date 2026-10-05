from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.orm import Session

from inventura.core.counting import DocumentStatus
from inventura.db.models import CountDocument


def found(client: TestClient, document: dict[str, Any], **body: object) -> Any:
    payload = {"lokacija": "B6-3-1", "sifra": "0000009", "presteta_kolicina": "2"} | body
    return client.post(f"/api/documents/{document['id']}/items", json=payload)


def close(session: Session, document: dict[str, Any]) -> None:
    session.execute(
        update(CountDocument)
        .where(CountDocument.id == document["id"])
        .values(status=DocumentStatus.ZAKLJUCEN)
    )


def test_found_goods_get_book_zero_and_current_price(
    client: TestClient, document: dict[str, Any]
) -> None:
    response = found(client, document, sarza="L5", stevec="Ana")

    assert response.status_code == 201, response.text
    item = response.json()
    assert (item["sifra"], item["lokacija"], item["sarza"], item["najdeno"]) == (
        "0000009",
        "B6-3-1",
        "L5",
        True,
    )
    assert (item["knjizena_kolicina"], item["presteta_kolicina"], item["cena_na_enoto"]) == (
        "0.000",
        "2.000",
        "40.00",
    )
    detail = client.get(f"/api/documents/{document['id']}").json()
    assert (detail["status"], detail["presteto"], detail["stevilo_postavk"]) == ("v_stetju", 1, 4)
    assert [i["lokacija"] for i in detail["postavke"]][-1] == "B6-3-1"


@pytest.mark.parametrize(
    ("body", "status", "message"),
    [
        ({"sifra": "7777777"}, 422, "not in the material master data"),
        ({"lokacija": "K1-01-01"}, 422, "not in rack B6"),
        ({"lokacija": "B6"}, 422, "invalid location"),
        ({"presteta_kolicina": "0"}, 422, "greater than 0"),
        ({"presteta_kolicina": "1.5"}, 422, "whole number"),
        ({"lokacija": "B6-01-01", "sifra": "0000001"}, 409, "already on the document"),
    ],
)
def test_found_goods_rules(
    client: TestClient, document: dict[str, Any], body: dict[str, str], status: int, message: str
) -> None:
    response = found(client, document, **body)
    assert response.status_code == status
    assert message in response.json()["detail"]


def test_found_goods_on_closed_document(
    client: TestClient, session: Session, document: dict[str, Any]
) -> None:
    close(session, document)
    assert found(client, document).status_code == 409


def test_only_found_goods_can_be_deleted(client: TestClient, document: dict[str, Any]) -> None:
    item_id = found(client, document).json()["id"]
    book_item = document["postavke"][0]["id"]

    assert client.delete(f"/api/items/{book_item}").status_code == 409
    assert client.delete(f"/api/items/{item_id}").status_code == 204
    assert client.delete(f"/api/items/{item_id}").status_code == 404
    assert client.get(f"/api/documents/{document['id']}").json()["stevilo_postavk"] == 3


def filled_sheet(client: TestClient, document: dict[str, Any], values: dict[str, object]) -> bytes:
    response = client.get(f"/api/documents/{document['id']}/count-sheet.xlsx")
    workbook = load_workbook(BytesIO(response.content))
    for cell, value in values.items():
        workbook["B6"][cell] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(client: TestClient, document: dict[str, Any], data: bytes, **form: str) -> Any:
    return client.post(
        f"/api/documents/{document['id']}/count-sheet",
        files={"file": ("popisni_list_B6.xlsx", data)},
        data=form,
    )


def test_upload_filled_sheet(client: TestClient, document: dict[str, Any]) -> None:
    vijak = document["postavke"][0]["id"]
    client.patch(f"/api/items/{vijak}", json={"presteta_kolicina": "100"})
    # Rows 6-8 are B6-1-1 vijak, B6-1-2 kabel, B6-2-1 olje; row 9 is written by hand.
    data = filled_sheet(
        client,
        document,
        {"G6": 118, "G7": "12,5", "B9": "B6-4-4", "C9": "0000009", "G9": 1},
    )

    response = upload(client, document, data, stevec="Bojan")

    assert response.status_code == 200, response.text
    assert response.json() == {"posodobljeno": 2, "prepisano": 1, "najdeno": 1, "brez_kolicine": 1}
    items = {
        (i["lokacija"], i["sifra"]): i
        for i in client.get(f"/api/documents/{document['id']}").json()["postavke"]
    }
    assert items[("B6-1-1", "0000001")]["presteta_kolicina"] == "118.000"
    assert items[("B6-1-1", "0000001")]["stevec"] == "Bojan"
    assert items[("B6-1-2", "0000002")]["presteta_kolicina"] == "12.500"
    assert items[("B6-2-1", "0000003")]["presteta_kolicina"] is None
    assert items[("B6-4-4", "0000009")]["najdeno"] is True


def test_upload_with_errors_stores_nothing(client: TestClient, document: dict[str, Any]) -> None:
    data = filled_sheet(client, document, {"G6": 5, "G7": "12,55"})

    response = upload(client, document, data)

    assert response.status_code == 422
    body = response.json()
    assert body["vrstic_z_napakami"] == 1
    assert body["napake"][0]["vrstica"] == 7
    assert client.get(f"/api/documents/{document['id']}").json()["presteto"] == 0


@pytest.mark.parametrize(
    ("data", "message"),
    [(b"not excel", "cannot read the file"), (None, "no worksheet named 'B6'")],
)
def test_upload_unreadable(
    client: TestClient, document: dict[str, Any], data: bytes | None, message: str
) -> None:
    if data is None:  # a workbook with two sheets, none of them B6
        workbook = load_workbook(
            BytesIO(client.get(f"/api/documents/{document['id']}/count-sheet.xlsx").content)
        )
        workbook["B6"].title = "X"
        workbook.create_sheet("Y")
        buffer = BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
    response = upload(client, document, data)
    assert response.status_code == 400
    assert message in response.json()["detail"]


def test_upload_to_closed_document(
    client: TestClient, session: Session, document: dict[str, Any]
) -> None:
    data = filled_sheet(client, document, {"G6": 5})
    close(session, document)
    assert upload(client, document, data).status_code == 409
