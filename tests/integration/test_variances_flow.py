"""Variances, recount rounds and the report, through the API and the pages.

The B6 document of the shared fixture has: 0000001 vijak 120 kos at 0,15 €, 0000002 kabel
12,5 m at 0,80 € and 0000003 olje 20 l (batch L01) at 3,50 €.
"""

from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from typer.testing import CliRunner

from inventura.cli import app as cli_app
from inventura.settings import Settings, get_settings


def ids(client: TestClient, document: dict[str, Any]) -> dict[str, int]:
    detail = client.get(f"/api/documents/{document['id']}").json()
    return {item["sifra"]: item["id"] for item in detail["postavke"]}


def count(client: TestClient, item_id: int, quantity: str) -> Any:
    return client.patch(f"/api/items/{item_id}", json={"presteta_kolicina": quantity})


def count_all(client: TestClient, document: dict[str, Any], olje: str = "5") -> None:
    items = ids(client, document)
    for sifra, quantity in (("0000001", "100"), ("0000002", "12.5"), ("0000003", olje)):
        assert count(client, items[sifra], quantity).status_code == 200


def finish(client: TestClient, document: dict[str, Any]) -> Any:
    return client.post(f"/api/documents/{document['id']}/finish-round")


def test_variances_are_computed_from_the_counts(
    client: TestClient, document: dict[str, Any]
) -> None:
    count_all(client, document)
    body = client.get(f"/api/documents/{document['id']}/variances").json()

    rows = {row["sifra"]: row for row in body["postavke"]}
    assert rows["0000001"]["razlika_kolicina"] == "-20.000"
    assert rows["0000001"]["razlika_vrednost"] == "-3.00"
    assert rows["0000001"]["ponovno_stetje"] is True  # any difference, whatever its value
    assert rows["0000002"]["ponovno_stetje"] is False
    assert rows["0000002"]["razlika_vrednost"] == "0.00"
    assert rows["0000003"]["razlika_vrednost"] == "-52.50"
    assert rows["0000003"]["odstopanje_odstotek"] == "-75.0"
    assert rows["0000003"]["ponovno_stetje"] is True
    assert body["povzetek"] == {
        "postavk": 3,
        "presteto": 3,
        "nepresteto": 0,
        "z_razliko": 2,
        "za_ponovno_stetje": 2,
        "visek": "0.00",
        "manjko": "-55.50",
        "neto": "-55.50",
        "absolutno": "55.50",
    }
    assert body["trenutni_krog"] == 1


def test_round_cannot_finish_with_uncounted_items(
    client: TestClient, document: dict[str, Any]
) -> None:
    response = finish(client, document)
    assert response.status_code == 409
    assert response.json()["detail"] == "krog še ni končan, nepreštetih postavk: 3"


def test_recount_round_and_closing(client: TestClient, document: dict[str, Any]) -> None:
    count_all(client, document)
    first_round = ids(client, document)

    result = finish(client, document)
    assert result.status_code == 200
    assert result.json() == {
        "status": "ponovno_stetje",
        "zakljuceni_krog": 1,
        "za_ponovno_stetje": 2,
    }

    # The screws and the oil go to the second round; their new rows start uncounted.
    detail = client.get(f"/api/documents/{document['id']}").json()
    rounds = {i["sifra"]: (i["krog"], i["presteta_kolicina"]) for i in detail["postavke"]}
    assert rounds == {
        "0000001": (2, None),
        "0000002": (1, "12.500"),
        "0000003": (2, None),
    }
    assert detail["status"] == "ponovno_stetje"

    # Items outside the recount are locked; the old row of the oil is history.
    locked = count(client, first_round["0000002"], "12")
    assert locked.status_code == 409
    assert locked.json()["detail"] == "v ponovnem štetju se štejejo samo postavke za ponovno štetje"
    assert count(client, first_round["0000003"], "6").status_code == 409

    second = ids(client, document)["0000003"]
    assert count(client, ids(client, document)["0000001"], "120").status_code == 200
    assert count(client, second, "19").status_code == 200
    variances = client.get(f"/api/documents/{document['id']}/variances").json()
    assert {r["sifra"]: r["razlika_vrednost"] for r in variances["postavke"]}["0000003"] == "-3.50"

    closed = finish(client, document)
    assert closed.json() == {"status": "zakljucen", "zakljuceni_krog": 2, "za_ponovno_stetje": 0}
    assert client.get(f"/api/documents/{document['id']}").json()["zakljucen_ob"] is not None
    assert finish(client, document).status_code == 409
    assert count(client, second, "20").status_code == 409


def test_last_round_accepts_the_count_even_over_the_threshold(
    client: TestClient, document: dict[str, Any]
) -> None:
    count_all(client, document)
    finish(client, document)
    count(client, ids(client, document)["0000001"], "100")  # the same shortages again
    count(client, ids(client, document)["0000003"], "5")
    assert finish(client, document).json()["status"] == "zakljucen"


def test_document_without_variances_closes_at_once(
    client: TestClient, document: dict[str, Any]
) -> None:
    items = ids(client, document)
    for sifra, quantity in (("0000001", "120"), ("0000002", "12.5"), ("0000003", "20")):
        count(client, items[sifra], quantity)
    assert finish(client, document).json() == {
        "status": "zakljucen",
        "zakljuceni_krog": 1,
        "za_ponovno_stetje": 0,
    }


def test_found_goods_always_go_to_recount(client: TestClient, document: dict[str, Any]) -> None:
    items = ids(client, document)
    for sifra, quantity in (("0000001", "120"), ("0000002", "12.5"), ("0000003", "20")):
        count(client, items[sifra], quantity)
    found = client.post(
        f"/api/documents/{document['id']}/items",
        json={"lokacija": "B6-3-1", "sifra": "0000009", "presteta_kolicina": "2"},  # 2 x 40 €
    )
    assert found.status_code == 201
    variances = client.get(f"/api/documents/{document['id']}/variances").json()
    row = next(r for r in variances["postavke"] if r["najdeno"])
    assert (row["razlika_vrednost"], row["odstopanje_odstotek"], row["ponovno_stetje"]) == (
        "80.00",
        None,
        True,
    )


def test_recount_sheet_and_page_show_only_recount_items(
    client: TestClient, document: dict[str, Any]
) -> None:
    count_all(client, document)
    finish(client, document)

    sheet = client.get(f"/api/documents/{document['id']}/count-sheet.xlsx")
    workbook = load_workbook(BytesIO(sheet.content))
    b6 = workbook["B6"]
    assert b6["A1"].value == "Popisni list – regal B6 – ponovno štetje (krog 2)"
    assert [b6[f"C{r}"].value for r in range(6, 9)] == ["0000001", "0000003", None]

    page = client.get(f"/documents/{document['id']}").text
    assert "Ponovno štetje (krog 2)" in page
    assert page.count('class="qty"') == 2
    assert "0 / 2" in page

    # A filled sheet may not count items outside the recount.
    b6["G6"] = 110
    b6["B8"], b6["C8"], b6["G8"] = "B6-1-2", "0000002", 12
    buffer = BytesIO()
    workbook.save(buffer)
    upload = client.post(
        f"/api/documents/{document['id']}/count-sheet",
        files={"file": ("b6.xlsx", buffer.getvalue())},
    )
    assert upload.status_code == 422
    assert upload.json()["napake"][0]["sporocilo"] == "postavka ni v ponovnem štetju"


def test_variances_page_and_finish_from_the_page(
    client: TestClient, document: dict[str, Any]
) -> None:
    count_all(client, document)
    page = client.get(f"/documents/{document['id']}/variances").text
    assert "Zaključi krog in pošlji 2 v ponovno štetje" in page
    assert "-52,50 €" in page
    assert "-75,0 %" in page

    response = client.post(f"/documents/{document['id']}/finish-round", follow_redirects=False)
    assert response.status_code == 303
    assert "Ponovno štetje" in client.get(f"/documents/{document['id']}/variances").text


def test_report(client: TestClient, document: dict[str, Any]) -> None:
    count_all(client, document)
    snapshot_id = document["snapshot_id"]
    response = client.get(f"/api/snapshots/{snapshot_id}/report.xlsx")
    assert response.status_code == 200
    assert f"porocilo_razlik_uvoz_{snapshot_id}.xlsx" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content))
    summary = list(workbook["Povzetek"].iter_rows(min_row=6, values_only=True))
    racks = {row[1]: row for row in summary}
    assert racks["B6"][3:] == (3, 3, 0, 2, 0, -55.5, -55.5, 55.5)
    assert racks["K1"][3:6] == (1, 0, 1)
    variances = list(workbook["Razlike"].iter_rows(min_row=2, values_only=True))
    assert [(row[2], row[10]) for row in variances] == [("0000001", -3), ("0000003", -52.5)]


def test_report_needs_documents(client: TestClient) -> None:
    data = "Šifra materiala;Opis materiala;ME;Lokacija;Zaloga;Cena na enoto\n1;V;kos;B1-1-1;1;1\n"
    snapshot = client.post("/api/snapshots", files={"file": ("s.csv", data.encode())}).json()
    response = client.get(f"/api/snapshots/{snapshot['id']}/report.xlsx")
    assert response.status_code == 409
    assert client.get("/api/snapshots/999999/report.xlsx").status_code == 404


def test_cli_report_for_unknown_import(
    monkeypatch: pytest.MonkeyPatch, test_settings: Settings
) -> None:
    monkeypatch.setenv("INVENTURA_DATABASE_URL", test_settings.database_url)
    get_settings.cache_clear()
    try:
        result = CliRunner().invoke(cli_app, ["report", "999999"])
    finally:
        get_settings.cache_clear()
    assert result.exit_code == 2
    assert "Poročila ni mogoče pripraviti: uvoz 999999 ne obstaja" in result.output
