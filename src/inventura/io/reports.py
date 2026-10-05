"""Excel report of an inventory: summary by rack, variances, and items not counted.

Values are written as numbers (not formulas), so every viewer shows them without
recalculating. Conditional formatting colours shortages red and surpluses green and
marks the items that went to a recount.
"""

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from inventura.core.variance import Variance, VarianceRules, VarianceSummary, summarize
from inventura.io.formats import (
    MONEY_NUMBER_FORMAT,
    format_date,
    format_decimal,
    quantity_number_format,
)

PERCENT_FORMAT = "0.0%"
SIGNED_MONEY_FORMAT = "+#,##0.00;-#,##0.00;0.00"

_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
_RED_FILL = PatternFill("solid", fgColor="F8D7DA")
_GREEN_FILL = PatternFill("solid", fgColor="D4EDDA")
_ORANGE_FILL = PatternFill("solid", fgColor="FFE5B4")
_RED_FONT = Font(color="9C0006")
_GREEN_FONT = Font(color="006100")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(bottom=_THIN)


@dataclass(frozen=True, slots=True)
class ReportLine:
    lokacija: str
    sifra: str
    opis: str
    sarza: str | None
    merska_enota: str
    krog: int
    najdeno: bool
    variance: Variance | None  # None: not counted


@dataclass(frozen=True, slots=True)
class ReportRack:
    number: int
    rack: str
    status: str
    lines: list[ReportLine]

    @property
    def summary(self) -> VarianceSummary:
        return summarize(line.variance for line in self.lines)


@dataclass(frozen=True, slots=True)
class ReportHeader:
    snapshot: str  # e.g. "Uvoz 3 · zaloga.csv"
    created: date
    rules: VarianceRules


def write_variance_report(
    racks: Sequence[ReportRack], header: ReportHeader, target: str | BinaryIO
) -> None:
    workbook = Workbook()
    for default in workbook.worksheets:
        workbook.remove(default)
    _summary(workbook.create_sheet("Povzetek"), racks, header)
    _variances(workbook.create_sheet("Razlike"), racks, header.rules)
    _uncounted(workbook.create_sheet("Neprešteto"), racks)
    workbook.save(target)


def rules_text(rules: VarianceRules) -> str:
    return (
        f"Ponovno štetje: (|vrednost razlike| ≥ {format_decimal(rules.min_value, 2)} € in "
        f"|odstopanje| ≥ {format_decimal(rules.min_percent, 0)} %) ali "
        f"|vrednost razlike| ≥ {format_decimal(rules.always_value, 2)} €; "
        f"največ {rules.max_rounds} kroga štetja."
    )


def _summary(sheet: Worksheet, racks: Sequence[ReportRack], header: ReportHeader) -> None:
    sheet["A1"] = "Inventura – poročilo o razlikah"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"{header.snapshot} · pripravljeno {format_date(header.created)}"
    sheet["A3"] = rules_text(header.rules)

    columns = [
        ("Št.", 6),
        ("Regal", 9),
        ("Status", 16),
        ("Postavk", 10),
        ("Prešteto", 10),
        ("Neprešteto", 12),
        ("Z razliko", 11),
        ("Ponovno štetje", 15),
        ("Višek (€)", 14),
        ("Manjko (€)", 14),
        ("Neto razlika (€)", 17),
        ("Absolutna razlika (€)", 20),
    ]
    first = 5
    _header_row(sheet, first, columns)
    row = first
    for rack in racks:
        row += 1
        _summary_row(sheet, row, [rack.number, rack.rack, rack.status], rack.summary)
    total = summarize(line.variance for rack in racks for line in rack.lines)
    row += 1
    _summary_row(sheet, row, ["", "Skupaj", ""], total)
    for cell in sheet[row]:
        cell.font = Font(bold=True)
        cell.fill = _TOTAL_FILL

    last = row
    net = f"K{first + 1}:K{last}"
    sheet.conditional_formatting.add(
        net, _cell_rule(operator="lessThan", formula=["0"], font=_RED_FONT)
    )
    sheet.conditional_formatting.add(
        net, _cell_rule(operator="greaterThan", formula=["0"], font=_GREEN_FONT)
    )
    sheet.conditional_formatting.add(
        f"F{first + 1}:F{last}",
        _cell_rule(operator="greaterThan", formula=["0"], fill=_ORANGE_FILL),
    )
    sheet.conditional_formatting.add(
        f"H{first + 1}:H{last}", _cell_rule(operator="greaterThan", formula=["0"], fill=_RED_FILL)
    )
    sheet.freeze_panes = f"A{first + 1}"
    _print_setup(sheet, first, landscape=True)


def _summary_row(sheet: Worksheet, row: int, labels: list[str | int], s: VarianceSummary) -> None:
    values: list[str | int | Decimal] = [
        *labels,
        s.items,
        s.counted,
        s.uncounted,
        s.with_difference,
        s.recount,
        s.surplus_value,
        s.shortage_value,
        s.net_value,
        s.absolute_value,
    ]
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column, value)
        cell.border = _BORDER
        if column >= 9:
            cell.number_format = SIGNED_MONEY_FORMAT if column == 11 else MONEY_NUMBER_FORMAT


def _variances(sheet: Worksheet, racks: Sequence[ReportRack], rules: VarianceRules) -> None:
    columns = [
        ("Regal", 8),
        ("Lokacija", 11),
        ("Šifra", 11),
        ("Opis", 40),
        ("Šarža", 11),
        ("ME", 5),
        ("Knjižna", 11),
        ("Prešteta", 11),
        ("Razlika", 11),
        ("Cena (€)", 11),
        ("Vrednost razlike (€)", 19),
        ("Odstopanje", 12),
        ("Krog", 6),
        ("Ponovno štetje", 15),
        ("Najdeno", 9),
    ]
    _header_row(sheet, 1, columns)
    row = 1
    for rack in racks:
        for line in rack.lines:
            variance = line.variance
            if variance is None or not variance.has_difference:
                continue
            row += 1
            unit_format = quantity_number_format(line.merska_enota)
            values: list[str | int | Decimal | None] = [
                rack.rack,
                line.lokacija,
                line.sifra,
                line.opis,
                line.sarza or "",
                line.merska_enota,
                variance.book,
                variance.counted,
                variance.quantity,
                variance.unit_price,
                variance.value,
                None if variance.percent is None else variance.percent / 100,
                line.krog,
                "da" if variance.recount else "ne",
                "da" if line.najdeno else "",
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row, column, value)
                cell.border = _BORDER
            for column in (7, 8, 9):
                sheet.cell(row, column).number_format = unit_format
            sheet.cell(row, 10).number_format = MONEY_NUMBER_FORMAT
            sheet.cell(row, 11).number_format = SIGNED_MONEY_FORMAT
            sheet.cell(row, 12).number_format = PERCENT_FORMAT

    if row == 1:
        sheet.cell(2, 1, "Ni razlik med knjižnim in preštetim stanjem.")
        return
    values_range = f"K2:K{row}"
    sheet.conditional_formatting.add(
        values_range, _cell_rule(operator="lessThan", formula=["0"], fill=_RED_FILL, font=_RED_FONT)
    )
    sheet.conditional_formatting.add(
        values_range,
        _cell_rule(operator="greaterThan", formula=["0"], fill=_GREEN_FILL, font=_GREEN_FONT),
    )
    # Variances that alone force a recount stand out across the whole row.
    always = str(rules.always_value)
    sheet.conditional_formatting.add(
        f"A2:O{row}", _formula_rule(formula=[f"ABS($K2)>={always}"], font=Font(bold=True))
    )
    sheet.conditional_formatting.add(
        f"N2:N{row}", _cell_rule(operator="equal", formula=['"da"'], fill=_ORANGE_FILL)
    )
    sheet.auto_filter.ref = f"A1:O{row}"
    sheet.freeze_panes = "A2"
    _print_setup(sheet, 1, landscape=True)


def _uncounted(sheet: Worksheet, racks: Sequence[ReportRack]) -> None:
    columns = [
        ("Regal", 8),
        ("Lokacija", 11),
        ("Šifra", 11),
        ("Opis", 40),
        ("Šarža", 11),
        ("ME", 5),
    ]
    _header_row(sheet, 1, columns)
    row = 1
    for rack in racks:
        for line in rack.lines:
            if line.variance is not None:
                continue
            row += 1
            values = [
                rack.rack,
                line.lokacija,
                line.sifra,
                line.opis,
                line.sarza or "",
                line.merska_enota,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row, column, value).border = _BORDER
    if row == 1:
        sheet.cell(2, 1, "Vse postavke so preštete.")
        return
    sheet.auto_filter.ref = f"A1:F{row}"
    sheet.freeze_panes = "A2"
    _print_setup(sheet, 1, landscape=False)


def _header_row(sheet: Worksheet, row: int, columns: Sequence[tuple[str, float]]) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row, index, title)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width


def _print_setup(sheet: Worksheet, header_row: int, landscape: bool) -> None:
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


# openpyxl's rule factories are untyped; these wrappers keep the rest of the module checked.
def _cell_rule(
    operator: str, formula: list[str], fill: PatternFill | None = None, font: Font | None = None
) -> Rule:
    rule: Rule = CellIsRule(operator=operator, formula=formula, fill=fill, font=font)  # type: ignore[no-untyped-call]
    return rule


def _formula_rule(formula: list[str], font: Font | None = None) -> Rule:
    rule: Rule = FormulaRule(formula=formula, font=font)  # type: ignore[no-untyped-call]
    return rule
