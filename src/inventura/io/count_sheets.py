"""Count sheets for rack documents: an Excel workbook and a printable HTML page.

Both list the items of each rack in walking order with an empty column for the counted
quantity. Book quantities are hidden by default (blind count), so the counter records
what is on the shelf rather than confirming the expected number.
"""

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import BinaryIO, Literal

from jinja2 import Environment, PackageLoader, StrictUndefined, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.header_footer import HeaderFooterItem
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from inventura.core.documents import RackDocument
from inventura.core.stock import StockRow
from inventura.core.units import quantity_decimals
from inventura.io.formats import format_date, format_quantity, quantity_number_format


@dataclass(frozen=True, slots=True)
class SheetOptions:
    created: date
    total_documents: int
    show_book_quantity: bool = False
    count_round: int = 1  # above 1: the sheet lists only the items sent to a recount


@dataclass(frozen=True, slots=True)
class _Column:
    title: str
    width: float


LOCATION_COLUMN_TITLE = "Lokacija"
CODE_COLUMN_TITLE = "Šifra"
BATCH_COLUMN_TITLE = "Šarža"
COUNT_COLUMN_TITLE = "Prešteta količina"
HEADER_ROW = 5

_THIN = Side(style="thin", color="808080")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_LEVEL_BORDER = Border(left=_THIN, right=_THIN, top=Side(style="medium"), bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_COUNT_FILL = PatternFill("solid", fgColor="FFF7D6")


def columns(show_book_quantity: bool) -> list[_Column]:
    result = [
        _Column("Zap. št.", 7),
        _Column(LOCATION_COLUMN_TITLE, 11),
        _Column(CODE_COLUMN_TITLE, 11),
        _Column("Opis", 44),
        _Column(BATCH_COLUMN_TITLE, 11),
        _Column("ME", 5),
    ]
    if show_book_quantity:
        result.append(_Column("Knjižna količina", 11))
    result += [_Column(COUNT_COLUMN_TITLE, 13), _Column("Opomba", 22)]
    return result


def document_title(document: RackDocument, count_round: int = 1) -> str:
    title = f"Popisni list – regal {document.rack}"
    return title if count_round == 1 else f"{title} – ponovno štetje (krog {count_round})"


def write_count_sheets_xlsx(
    documents: Sequence[RackDocument], target: Path | BinaryIO, options: SheetOptions
) -> None:
    """One worksheet per rack, set up to print on A4 with the header repeated on every page."""
    if not documents:
        raise ValueError("no documents to write")
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    workbook.remove(first)
    for document in documents:
        _write_sheet(workbook.create_sheet(document.rack), document, options)
    workbook.save(target)


def render_count_sheets_html(documents: Sequence[RackDocument], options: SheetOptions) -> str:
    """A single page with every document; each starts on a new sheet of paper when printed."""
    environment = Environment(
        loader=PackageLoader("inventura", "web/templates"),
        autoescape=select_autoescape(),
        undefined=StrictUndefined,
        trim_blocks=True,
        lstrip_blocks=True,
    )
    environment.filters["quantity"] = lambda item: format_quantity(item.kolicina, item.merska_enota)
    template = environment.get_template("count_sheets.html")
    return template.render(
        documents=documents,
        columns=columns(options.show_book_quantity),
        show_book_quantity=options.show_book_quantity,
        created=format_date(options.created),
        total_documents=options.total_documents,
        title=lambda document: document_title(document, options.count_round),
        new_level=_new_level_flags,
    )


def write_count_sheets_html(
    documents: Sequence[RackDocument], path: Path, options: SheetOptions
) -> None:
    path.write_text(render_count_sheets_html(documents, options), encoding="utf-8")


def _new_level_flags(items: Sequence[StockRow]) -> list[bool]:
    """True for the first item of each level after the first, to draw a separator line."""
    flags = []
    previous: int | None = None
    for item in items:
        flags.append(previous is not None and item.lokacija.level != previous)
        previous = item.lokacija.level
    return flags


def _write_sheet(sheet: Worksheet, document: RackDocument, options: SheetOptions) -> None:
    cols = columns(options.show_book_quantity)
    count_col = next(i for i, c in enumerate(cols, start=1) if c.title == COUNT_COLUMN_TITLE)

    sheet["A1"] = document_title(document, options.count_round)
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        f"Dokument {document.number} od {options.total_documents} · "
        f"Datum: {format_date(options.created)} · "
        f"Postavk: {len(document.items)} · Lokacij: {len(document.locations)}"
    )
    sheet["A3"] = "Števec: ______________________      Podpis: ______________________"

    for index, column in enumerate(cols, start=1):
        cell = sheet.cell(HEADER_ROW, index, column.title)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[cell.column_letter].width = column.width

    whole = _count_validation("whole", "Vnesi celo število, 0 ali več.")
    decimal = _count_validation("decimal", "Vnesi število, 0 ali več.")

    for offset, (item, new_level) in enumerate(
        zip(document.items, _new_level_flags(document.items), strict=True)
    ):
        row = HEADER_ROW + 1 + offset
        values: list[str | int | Decimal | None] = [
            offset + 1,
            str(item.lokacija),
            item.sifra,
            item.opis,
            item.sarza or "",
            item.merska_enota,
        ]
        if options.show_book_quantity:
            values.append(item.kolicina)
        values += [None, None]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row, index, value)
            cell.border = _LEVEL_BORDER if new_level else _BORDER
        quantity_format = quantity_number_format(item.merska_enota)
        if options.show_book_quantity:
            sheet.cell(row, count_col - 1).number_format = quantity_format
        count_cell = sheet.cell(row, count_col)
        count_cell.fill = _COUNT_FILL
        count_cell.number_format = quantity_format
        (decimal if quantity_decimals(item.merska_enota) else whole).add(count_cell)

    for validation in (whole, decimal):
        if validation.sqref.ranges:
            sheet.add_data_validation(validation)

    sheet.freeze_panes = f"A{HEADER_ROW + 1}"
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    sheet.print_area = f"A1:{sheet.cell(sheet.max_row, len(cols)).coordinate}"
    sheet.page_setup.orientation = "landscape" if options.show_book_quantity else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    footer = HeaderFooterItem()
    footer.left.text = f"Regal {document.rack}"
    footer.right.text = "Stran &P od &N"
    sheet.oddFooter = footer


def _count_validation(kind: Literal["whole", "decimal"], message: str) -> DataValidation:
    validation = DataValidation(type=kind, operator="greaterThanOrEqual", formula1="0")
    validation.error = message
    validation.showErrorMessage = True
    return validation
