"""Read a filled count sheet (the Excel written by count_sheets.py) back in.

The worksheet is the one named after the rack, or the only one in the workbook. The header
row is found by its "Prešteta količina" cell, so rows added above it do not matter.
Counters may write found goods into empty rows at the end of the table.
"""

import zipfile
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from inventura.core.count_sheet_import import SheetRow
from inventura.io.count_sheets import (
    BATCH_COLUMN_TITLE,
    CODE_COLUMN_TITLE,
    COUNT_COLUMN_TITLE,
    LOCATION_COLUMN_TITLE,
)

HEADER_SEARCH_ROWS = 20


class CountSheetError(ValueError):
    """The file is not a readable count sheet for this rack."""


def read_count_sheet(data: bytes, rack: str) -> list[SheetRow]:
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise CountSheetError(f"datoteke ni mogoče prebrati kot Excel: {exc}") from exc
    sheet = _sheet_for_rack(workbook.sheetnames, rack)
    worksheet = workbook[sheet]
    if not isinstance(worksheet, Worksheet):  # e.g. a chart sheet
        raise CountSheetError(f"{sheet!r} ni delovni list")
    return _rows(worksheet)


def _sheet_for_rack(names: list[str], rack: str) -> str:
    for name in names:
        if name.strip().casefold() == rack.casefold():
            return name
    if len(names) == 1:
        return names[0]
    raise CountSheetError(f"v datoteki ni lista z imenom {rack!r}")


def _rows(worksheet: Worksheet) -> list[SheetRow]:
    rows = worksheet.iter_rows(values_only=True)
    header_row, columns = 0, {}
    for number, values in enumerate(rows, start=1):
        titles = {_title(v): i for i, v in enumerate(values) if v is not None}
        if _title(COUNT_COLUMN_TITLE) in titles:
            header_row, columns = number, titles
            break
        if number >= HEADER_SEARCH_ROWS:
            break
    if not header_row:
        raise CountSheetError(f"ni vrstice z naslovom stolpca {COUNT_COLUMN_TITLE!r}")

    required = (LOCATION_COLUMN_TITLE, CODE_COLUMN_TITLE, COUNT_COLUMN_TITLE)
    missing = [title for title in required if _title(title) not in columns]
    if missing:
        raise CountSheetError(f"manjkajo stolpci: {', '.join(missing)}")

    def cell(values: tuple[object, ...], title: str) -> object:
        index = columns.get(_title(title))
        return values[index] if index is not None and index < len(values) else None

    return [
        SheetRow(
            row_number=number,
            lokacija=cell(values, LOCATION_COLUMN_TITLE),
            sifra=cell(values, CODE_COLUMN_TITLE),
            sarza=cell(values, BATCH_COLUMN_TITLE),
            kolicina=cell(values, COUNT_COLUMN_TITLE),
        )
        for number, values in enumerate(rows, start=header_row + 1)
    ]


def _title(value: object) -> str:
    return " ".join(str(value).split()).casefold()
